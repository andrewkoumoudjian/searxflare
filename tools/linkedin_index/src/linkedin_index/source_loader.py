from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence
from urllib.parse import unquote, urlsplit

from .canonical import canonical_linkedin_url
from .normalize import NormalizedEntity, SourceRecord, collapse_person_aliases, merge_records


SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xls", ".numbers"}
LINKEDIN_RE = re.compile(
    r"(?:(?:https?://)?(?:[a-z]{2,5}\.)?(?:www\.)?linkedin\.com/"
    r"(?:in|company|school|pulse|posts|showcase)/[^\s,;\"'<>]+)",
    re.IGNORECASE,
)


@dataclass
class LoadReport:
    files_total: int = 0
    files_read: int = 0
    files_failed: int = 0
    rows_read: int = 0
    source_records: int = 0
    unique_entities: int = 0
    people: int = 0
    companies: int = 0
    schools: int = 0
    posts: int = 0
    records_with_email: int = 0
    records_with_phone: int = 0
    records_with_both: int = 0
    duplicates_merged: int = 0
    conflicting_fields: int = 0
    errors: list[dict[str, str]] = field(default_factory=list)
    per_file: dict[str, dict[str, Any]] = field(default_factory=dict)


@dataclass
class LoadedSources:
    records: list[SourceRecord]
    report: LoadReport


@dataclass
class NormalizedSources:
    entities: dict[str, NormalizedEntity]
    report: LoadReport

    def to_jsonable(self) -> dict[str, Any]:
        return {
            "report": asdict(self.report),
            "entities": {
                url: {
                    "url": entity.url,
                    "entity_type": entity.entity_type,
                    "fields": {
                        name: [
                            {
                                "value": _json_safe(value.value),
                                "source_file": value.source_file,
                                "row_number": value.row_number,
                            }
                            for value in values
                        ]
                        for name, values in entity.fields.items()
                    },
                    "source_records": [
                        {
                            "source_file": record.source_file,
                            "row_number": record.row_number,
                            "linkedin_url": record.linkedin_url,
                        }
                        for record in entity.source_records
                    ],
                }
                for url, entity in self.entities.items()
            },
        }


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        return {str(key): _json_safe(child) for key, child in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(child) for child in value]
    return str(value)


def _clean_header(value: Any, index: int, seen: dict[str, int]) -> str:
    base = str(value).strip() if value not in (None, "") else f"column_{index + 1}"
    count = seen.get(base, 0)
    seen[base] = count + 1
    return base if count == 0 else f"{base}__{count + 1}"


def _rows_from_matrix(rows: Iterable[Sequence[Any]]) -> Iterator[tuple[int, dict[str, Any]]]:
    iterator = iter(rows)
    headers: list[str] | None = None
    physical_row = 0
    for raw in iterator:
        physical_row += 1
        values = list(raw)
        if headers is None:
            if not any(value not in (None, "") for value in values):
                continue
            seen: dict[str, int] = {}
            headers = [_clean_header(value, index, seen) for index, value in enumerate(values)]
            continue
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        elif len(values) > len(headers):
            seen = {header: 1 for header in headers}
            headers.extend(
                _clean_header(None, index, seen)
                for index in range(len(headers), len(values))
            )
        row = {headers[index]: values[index] for index in range(len(headers))}
        if any(_nonempty(value) for value in row.values()):
            yield physical_row, row


def _read_csv(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        sample = handle.read(65536)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        yield from _rows_from_matrix(csv.reader(handle, dialect=dialect))


def _read_xlsx(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            yield from _rows_from_matrix(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def _read_xls(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    import xlrd

    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
    except xlrd.XLRDError:
        prefix = path.read_bytes()[:256].lstrip()
        if prefix.startswith(b"<?xml"):
            yield from _read_excel_2003_xml(path)
            return
        raise
    try:
        for sheet in workbook.sheets():
            rows = (
                [sheet.cell_value(row, column) for column in range(sheet.ncols)]
                for row in range(sheet.nrows)
            )
            yield from _rows_from_matrix(rows)
    finally:
        workbook.release_resources()


def _read_excel_2003_xml(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    namespace = "urn:schemas-microsoft-com:office:spreadsheet"
    ss_index = f"{{{namespace}}}Index"
    data_tag = f"{{{namespace}}}Data"
    row_tag = f"{{{namespace}}}Row"
    cell_tag = f"{{{namespace}}}Cell"

    def matrix_rows() -> Iterator[list[Any]]:
        for _event, element in ET.iterparse(path, events=("end",)):
            if element.tag != row_tag:
                continue
            values: list[Any] = []
            cursor = 1
            for cell in element.findall(cell_tag):
                index_raw = cell.attrib.get(ss_index)
                if index_raw:
                    target = max(1, int(index_raw))
                    while cursor < target:
                        values.append(None)
                        cursor += 1
                data = cell.find(data_tag)
                values.append(data.text if data is not None else None)
                cursor += 1
            yield values
            element.clear()

    yield from _rows_from_matrix(matrix_rows())


def _read_numbers(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    from numbers_parser import Document

    document = Document(path)
    row_offset = 0
    for sheet in document.sheets:
        for table in sheet.tables:
            rows = (
                [table.cell(row, column).value for column in range(table.num_cols)]
                for row in range(table.num_rows)
            )
            for row_number, record in _rows_from_matrix(rows):
                yield row_offset + row_number, record
            row_offset += table.num_rows + 1


def _read_file(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        yield from _read_csv(path)
    elif suffix == ".xlsx":
        yield from _read_xlsx(path)
    elif suffix == ".xls":
        yield from _read_xls(path)
    elif suffix == ".numbers":
        yield from _read_numbers(path)
    else:
        raise ValueError(f"unsupported source format: {path.suffix}")


def _nonempty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _normalized_key(key: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", key.casefold())


def _first_value(row: dict[str, Any], keys: set[str]) -> Any | None:
    for key, value in row.items():
        if _normalized_key(key) in keys and _nonempty(value):
            return value
    return None


def _canonical_aliases(row: dict[str, Any]) -> dict[str, Any]:
    aliases: dict[str, Any] = {}
    first = _first_value(row, {"firstname", "first"})
    last = _first_value(row, {"lastname", "last"})
    explicit_name = _first_value(
        row,
        {"name", "fullname", "contactname", "personname", "linkedinname", "profilename"},
    )
    if explicit_name:
        aliases["name"] = explicit_name
    elif first or last:
        aliases["name"] = " ".join(str(value).strip() for value in (first, last) if _nonempty(value))

    alias_keys: dict[str, set[str]] = {
        "title": {"title", "jobtitle", "currenttitle", "linkedinjobtitle", "headline"},
        "company": {"company", "companyname", "currentcompany", "companynameforemails"},
        "email": {"email", "primaryemail", "workemail", "personalemail"},
        "phone": {
            "phone", "phonenumber", "mobilephone", "workdirectphone", "workphone",
            "corporatephone", "homephone", "otherphone", "phonenumber2",
        },
        "location": {"location", "linkedinlocation", "linkedinjoblocation", "personlocation"},
        "industry": {"industry", "companyindustry"},
    }
    for alias, keys in alias_keys.items():
        value = _first_value(row, keys)
        if value is not None:
            aliases[alias] = value

    if "location" not in aliases:
        city = _first_value(row, {"city"})
        state = _first_value(row, {"state", "province"})
        country = _first_value(row, {"country"})
        parts = [str(value).strip() for value in (city, state, country) if _nonempty(value)]
        if parts:
            aliases["location"] = ", ".join(parts)
    return aliases


def extract_linkedin_urls(row: dict[str, Any]) -> list[str]:
    urls: list[str] = []
    seen: set[str] = set()
    for value in row.values():
        if not isinstance(value, str) or "linkedin.com/" not in value.casefold():
            continue
        for raw in LINKEDIN_RE.findall(value):
            candidate = raw.strip().rstrip("/.)]}>")
            if not candidate.lower().startswith(("http://", "https://")):
                candidate = "https://" + candidate
            canonical = canonical_linkedin_url(candidate)
            if canonical and canonical not in seen:
                seen.add(canonical)
                urls.append(canonical)
    return urls


def load_source_tree(root: str | Path) -> LoadedSources:
    root_path = Path(root)
    files = sorted(
        path for path in root_path.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )
    report = LoadReport(files_total=len(files))
    records: list[SourceRecord] = []

    for path in files:
        relative = str(path.relative_to(root_path))
        file_rows = 0
        file_records = 0
        try:
            for row_number, raw_row in _read_file(path):
                file_rows += 1
                row = {str(key): value for key, value in raw_row.items() if _nonempty(value)}
                if not row:
                    continue
                urls = extract_linkedin_urls(row)
                if not urls:
                    continue
                fields = dict(row)
                fields.update(_canonical_aliases(row))
                for url in urls:
                    records.append(
                        SourceRecord(
                            source_file=relative,
                            row_number=row_number,
                            linkedin_url=url,
                            fields=fields,
                        )
                    )
                    file_records += 1
            report.files_read += 1
            report.rows_read += file_rows
            report.per_file[relative] = {
                "rows_read": file_rows,
                "source_records": file_records,
                "status": "ok",
            }
        except Exception as exc:
            report.files_failed += 1
            report.errors.append({"file": str(path), "error": f"{type(exc).__name__}: {exc}"})
            report.per_file[relative] = {
                "rows_read": file_rows,
                "source_records": file_records,
                "status": "failed",
                "error": f"{type(exc).__name__}: {exc}",
            }
    report.source_records = len(records)
    return LoadedSources(records=records, report=report)


def _has_field(entity: NormalizedEntity, field_name: str) -> bool:
    return any(_nonempty(item.value) for item in entity.fields.get(field_name, []))


def normalize_sources(root: str | Path) -> NormalizedSources:
    loaded = load_source_tree(root)
    entities = collapse_person_aliases(merge_records(loaded.records))
    report = loaded.report
    report.unique_entities = len(entities)
    report.people = sum(entity.entity_type == "person" for entity in entities.values())
    report.companies = sum(entity.entity_type == "company" for entity in entities.values())
    report.schools = sum(entity.entity_type == "school" for entity in entities.values())
    report.posts = sum(entity.entity_type == "post" for entity in entities.values())
    report.records_with_email = sum(_has_field(entity, "email") for entity in entities.values())
    report.records_with_phone = sum(_has_field(entity, "phone") for entity in entities.values())
    report.records_with_both = sum(
        _has_field(entity, "email") and _has_field(entity, "phone")
        for entity in entities.values()
    )
    report.duplicates_merged = max(0, len(loaded.records) - len(entities))
    report.conflicting_fields = sum(
        1
        for entity in entities.values()
        for values in entity.fields.values()
        if len({json.dumps(_json_safe(item.value), sort_keys=True) for item in values}) > 1
    )
    return NormalizedSources(entities=entities, report=report)


def select_pilot_seeds(entities: dict[str, NormalizedEntity], limit: int = 500) -> list[str]:
    if limit < 1:
        return []

    def stable_identity(url: str, entity: NormalizedEntity) -> tuple[str, str]:
        for wanted in ("apollocontactid", "apollorecordid"):
            values: list[str] = []
            for field_name, field_values in entity.fields.items():
                if _normalized_key(field_name) != wanted:
                    continue
                for item in field_values:
                    if _nonempty(item.value):
                        value = str(item.value).strip().casefold()
                        if value not in values:
                            values.append(value)
            if len(values) == 1:
                return (wanted, values[0])
        return ("url", url)

    def normalized_tokens(value: str) -> list[str]:
        ascii_value = (
            unicodedata.normalize("NFKD", value)
            .encode("ascii", "ignore")
            .decode("ascii")
            .casefold()
        )
        return [token for token in re.split(r"[^a-z0-9]+", ascii_value) if token]

    def alias_preference(item: tuple[str, NormalizedEntity]) -> tuple[int, int, int, str]:
        url, entity = item
        slug = unquote(urlsplit(url).path.rsplit("/", 1)[-1])
        slug_tokens = set(normalized_tokens(slug))
        name_values = entity.fields.get("name", [])
        name = str(name_values[0].value) if name_values else ""
        overlap = sum(token in slug_tokens for token in normalized_tokens(name) if len(token) > 1)
        return (overlap, len(slug), len(entity.fields), url)

    def score(item: tuple[str, NormalizedEntity]) -> tuple[int, int, int, str]:
        url, entity = item
        has_email = int(_has_field(entity, "email"))
        has_phone = int(_has_field(entity, "phone"))
        richness = len(entity.fields)
        return (-has_email, -has_phone, -richness, url)

    identity_groups: dict[tuple[str, str], list[tuple[str, NormalizedEntity]]] = {}
    for url, entity in entities.items():
        if entity.entity_type != "person":
            continue
        identity_groups.setdefault(stable_identity(url, entity), []).append((url, entity))

    representatives = [max(group, key=alias_preference) for group in identity_groups.values()]
    people = sorted(representatives, key=score)
    return [url for url, _ in people[:limit]]


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Normalize LinkedIn people/company enrichment files")
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--pilot-size", type=int, default=500)
    args = parser.parse_args(argv)

    result = normalize_sources(args.root)
    args.out.mkdir(parents=True, exist_ok=True)
    normalized_path = args.out / "normalized.json"
    normalized_path.write_text(json.dumps(result.to_jsonable(), ensure_ascii=False), encoding="utf-8")
    seeds = select_pilot_seeds(result.entities, limit=args.pilot_size)
    seed_path = args.out / "pilot-seeds.json"
    seed_path.write_text(json.dumps(seeds, indent=2), encoding="utf-8")
    report_path = args.out / "normalization-report.json"
    report_path.write_text(json.dumps(asdict(result.report), indent=2, sort_keys=True), encoding="utf-8")
    print(json.dumps({
        "normalized": str(normalized_path),
        "seeds": str(seed_path),
        "report": str(report_path),
        "pilot_seed_count": len(seeds),
        **asdict(result.report),
    }, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
